"""Generate synthetic fixtures and real processing results in a NEW directory."""

import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from pypdf import PdfReader, PdfWriter

from cleanroom.core import clean_document
from cleanroom.pdf_io import rebuild_pdf


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    directory = args.output.resolve()
    directory.mkdir(parents=True, exist_ok=False)
    document = Document()
    document.add_heading("人工合成测试样本 / Synthetic fixture", level=1)
    document.add_paragraph("这是正常正文，不是已确认带 textGrain 的 GPT 输出。")
    document.add_paragraph("字符清理：零\u200b宽空格、词\u2060连接符、字节\ufeff序标记、软\u00ad连字符。")
    document.add_paragraph("正常连接符保留：👩\u200d💻 / فارسی\u200c。")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip((c for row in table.rows for c in row.cells), ["项目", "值", "内容", "不改写"]):
        cell.text = value
    document.sections[0].header.paragraphs[0].text = "页眉候选字符\u200b：应仅去字符并保留页眉。"
    document.core_properties.author = "SYNTHETIC_AUTHOR_FOR_TEST"
    document.core_properties.last_modified_by = "SYNTHETIC_EDITOR_FOR_TEST"
    docx_source = directory / "synthetic-input.docx"
    document.save(docx_source)

    txt_source = directory / "synthetic-input.txt"
    txt_source.write_text("人工合成样本。\nHello\u200b world\u2060. 内容\ufeff不\u00ad改写。\n"
                          "这只能验证指定字符，不是官方统计水印实验。", encoding="utf-8")

    pdf_base = directory / "synthetic-base.pdf"
    rebuild_pdf("人工合成 PDF 验证样本\nPDF 文本内容保留。Metadata sample 2026.\n"
                "This fixture is not certified textGrain watermarked output.", pdf_base)
    pdf_source = directory / "synthetic-input.pdf"
    writer = PdfWriter(clone_from=PdfReader(pdf_base))
    writer.add_metadata({"/Author": "SYNTHETIC_PDF_AUTHOR", "/Title": "SYNTHETIC_METADATA"})
    with pdf_source.open("wb") as stream:
        writer.write(stream)

    cases = [
        (docx_source, "docx-preserved.docx", "preserve"),
        (txt_source, "text-cleaned.txt", "preserve"),
        (txt_source, "text-rebuilt.docx", "rebuild"),
        (txt_source, "text-rebuilt.pdf", "rebuild"),
        (pdf_source, "pdf-metadata-cleaned.pdf", "preserve"),
        (pdf_source, "pdf-rebuilt.pdf", "rebuild"),
    ]
    results = []
    for source, output, mode in cases:
        report = clean_document(source, directory / output, mode=mode)
        results.append({
            "input": source.name, "output": output,
            "selected_characters_before": report["before"]["invisible_characters"]["selected_count"],
            "selected_characters_after": report["after"]["invisible_characters"]["selected_count"],
            "metadata_fields_before": report["before"]["metadata_fields"],
            "metadata_fields_after": report["after"]["metadata_fields"],
            "text_comparison_passed": report["verification"]["expected_text_matches"],
            "textGrain": "not_tested",
        })
    summary = {
        "fixture_origin": "人工合成测试数据；不是经 OpenAI 确认的 textGrain 水印样本。",
        "results": results,
        "official_detector_run": False,
        "conclusion": "只能确认列出的字符和元数据处理结果，不能证明 textGrain 不可识别。",
    }
    (directory / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
