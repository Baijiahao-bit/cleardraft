import json
import os
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from cleanroom import cli, core
from cleanroom.pdf_io import rebuild_pdf


@pytest.mark.skipif(os.name != "nt", reason="Windows path safety")
@pytest.mark.parametrize("name", ["source.txt:hidden.txt", "nul.txt", "trailing. /out.txt"])
def test_windows_special_output_paths_are_rejected(tmp_path, name):
    source = tmp_path / "source.txt"
    source.write_text("original", encoding="utf-8")
    with pytest.raises(ValueError, match="Windows"):
        core._validate_output(source, tmp_path / name, "preserve")


@pytest.mark.parametrize("suffix", [".docx", ".pdf", ".txt"])
def test_cross_format_rebuild_verifies_output(tmp_path, suffix):
    source = tmp_path / "source.txt"
    source.write_text("保留\u200b中文。\nKeep\u2060 English.", encoding="utf-8")
    output = tmp_path / ("new" + suffix)
    report = core.clean_document(source, output, mode="rebuild")
    assert report["verification"]["expected_text_matches"]
    assert report["after"]["invisible_characters"]["selected_count"] == 0
    assert report["after"]["metadata_fields"] == []
    assert report["openai_text_watermark"]["verdict"] == "unknown"


def test_docx_preserve_and_no_source_text_in_audit(tmp_path):
    source = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("BODY_SECRET\u200b")
    document.sections[0].header.paragraphs[0].text = "HEADER_SECRET\u200b"
    document.save(source)
    output = tmp_path / "b.docx"
    report = core.clean_document(source, output)
    assert report["before"]["extra_text_invisible_characters"]["selected_count"] == 1
    assert report["after"]["extra_text_invisible_characters"]["selected_count"] == 0
    assert "HEADER_SECRET" not in json.dumps(report)
    assert "BODY_SECRET" not in json.dumps(report)


def test_rebuild_refuses_complex_docx_without_explicit_flag(tmp_path):
    source = tmp_path / "a.docx"
    document = Document()
    document.add_paragraph("正文")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "表格"
    document.save(source)
    with pytest.raises(ValueError, match="allow-lossy"):
        core.clean_document(source, tmp_path / "b.txt", mode="rebuild")
    report = core.clean_document(source, tmp_path / "b.txt", mode="rebuild", allow_lossy=True)
    assert report["verification"]["expected_text_matches"]


def test_pdf_preserve_metadata_report(tmp_path):
    original = tmp_path / "a.pdf"
    rebuild_pdf("Sample PDF", original)
    source = tmp_path / "tagged.pdf"
    writer = PdfWriter(clone_from=PdfReader(original))
    writer.add_metadata({"/Author": "secret author"})
    writer.write(source)
    report = core.clean_document(source, tmp_path / "clean.pdf")
    assert report["before"]["metadata_fields"]
    assert not report["after"]["metadata_fields"]
    assert report["operation"]["method"] == "pdf_metadata_only"


@pytest.mark.parametrize("case", ["suffix", "parent", "too_large", "format_mismatch"])
def test_validation_edges(tmp_path, monkeypatch, case):
    source = tmp_path / ("a.exe" if case == "suffix" else "a.txt")
    source.write_text("test", encoding="utf-8")
    out = tmp_path / "b.txt"
    if case == "parent":
        out = tmp_path / "missing" / "b.txt"
    elif case == "too_large":
        monkeypatch.setattr(core, "MAX_BYTES", 1)
    elif case == "format_mismatch":
        out = tmp_path / "b.docx"
    with pytest.raises(ValueError):
        core.clean_document(source, out)


def test_content_loss_is_rejected_before_publishing(tmp_path, monkeypatch):
    source = tmp_path / "a.txt"
    source.write_text("Must stay complete", encoding="utf-8")
    output = tmp_path / "b.txt"
    monkeypatch.setattr(core, "_write_rebuilt", lambda text, path: path.write_text("lost content"))
    with pytest.raises(ValueError, match="复核失败"):
        core.clean_document(source, output, mode="rebuild")
    assert not output.exists()
    assert not list(tmp_path.glob("cleanroom-*"))


def test_concurrently_changed_source_aborts(tmp_path, monkeypatch):
    source = tmp_path / "a.txt"
    source.write_text("first", encoding="utf-8")
    normal = core._transform

    def change_source(*args):
        result = normal(*args)
        source.write_text("changed", encoding="utf-8")
        return result

    monkeypatch.setattr(core, "_transform", change_source)
    with pytest.raises(ValueError, match="源文件被其他程序修改"):
        core.clean_document(source, tmp_path / "b.txt")
    assert not (tmp_path / "b.txt").exists()


def test_report_publish_failure_rolls_back_its_own_output(tmp_path):
    stage = tmp_path / "stage.txt"
    stage.write_text("new", encoding="utf-8")
    output = tmp_path / "out.txt"
    audit = tmp_path / "audit.json"
    audit.write_text("existing report")
    with pytest.raises(FileExistsError):
        core._publish(stage, output, audit, {})
    assert not output.exists()
    assert audit.read_text() == "existing report"


def test_cli_direct_paths_and_errors(tmp_path, capsys):
    source = tmp_path / "a.txt"
    source.write_text("clean\u200b", encoding="utf-8")
    assert cli.main(["inspect", str(source)]) == 0
    assert '"selected_count": 1' in capsys.readouterr().out
    assert cli.main(["clean", str(source), "-o", str(tmp_path / "b.txt")]) == 0
    assert cli.main(["inspect", str(tmp_path / "missing")]) == 2
