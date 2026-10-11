"""DOCX regression and input safety tests."""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import struct

import pytest
from docx import Document
from lxml import etree

from cleanroom.docx_io import clean_docx, extract_docx, inspect_docx, rebuild_docx
from cleanroom import docx_io


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def rewrite_zip(source: Path, destination: Path, replacements: dict[str, bytes]) -> None:
    with ZipFile(source) as archive:
        entries = {info.filename: archive.read(info) for info in archive.infolist()}
    with ZipFile(destination, "w", ZIP_DEFLATED) as archive:
        for name, data in {**entries, **replacements}.items():
            archive.writestr(name, data)


@pytest.fixture
def sample(tmp_path: Path) -> Path:
    path = tmp_path / "input.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("中\u200b文").bold = True
    paragraph.add_run(" test").italic = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "左列"
    table.cell(0, 1).text = "右列"
    document.add_paragraph("末尾")
    document.sections[0].header.paragraphs[0].text = "页\u200b眉"
    document.core_properties.author = "Private Author"
    document.core_properties.last_modified_by = "Private Editor"
    document.save(path)
    return path


def test_extract_preserves_table_document_order(sample: Path) -> None:
    assert extract_docx(sample) == "中\u200b文 test\n左列\n右列\n末尾"


def test_inspection_excludes_values_and_includes_nonbody_text(sample: Path) -> None:
    report = inspect_docx(sample)
    assert "core:creator" in report["metadata_fields"]
    assert "Private Author" not in repr(report)
    assert "页\u200b眉" in report["extra_text"]


def test_clean_preserves_styles_tables_headers_and_removes_metadata(sample: Path, tmp_path: Path) -> None:
    output = tmp_path / "output.docx"
    result = clean_docx(sample, output, lambda text: text.replace("\u200b", ""))
    assert extract_docx(output) == "中文 test\n左列\n右列\n末尾"
    reopened = Document(output)
    assert reopened.paragraphs[0].runs[0].bold
    assert reopened.paragraphs[0].runs[1].italic
    assert reopened.sections[0].header.paragraphs[0].text == "页眉"
    assert len(reopened.tables) == 1
    assert inspect_docx(output)["metadata_fields"] == []
    assert result["changed_text_nodes"] == 2
    assert "extra_text" not in result
    with ZipFile(output) as archive:
        assert not any(name.startswith("docProps/") for name in archive.namelist())
        assert b"docProps" not in archive.read("_rels/.rels")
        assert b"docProps" not in archive.read("[Content_Types].xml")


def test_rebuild_creates_plain_document_and_page_breaks(tmp_path: Path) -> None:
    path = tmp_path / "new.docx"
    rebuild_docx("第一行\n第二行\f新页", path)
    document = Document(path)
    assert "第一行" in extract_docx(path)
    assert "新页" in extract_docx(path)
    assert not document.core_properties.author
    assert inspect_docx(path)["metadata_fields"] == []
    with ZipFile(path) as archive:
        assert b'w:type="page"' in archive.read("word/document.xml")


def test_clean_removes_rsid_but_reports_preserved_markers(sample: Path, tmp_path: Path) -> None:
    marked = tmp_path / "marked.docx"
    with ZipFile(sample) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    root.set(f"{{{W}}}rsidR", "00112233")
    change = etree.SubElement(root.find(f"{{{W}}}body"), f"{{{W}}}ins")
    change.set(f"{{{W}}}author", "Writer")
    rewrite_zip(sample, marked, {
        "word/document.xml": etree.tostring(root),
        "word/comments.xml": f'<w:comments xmlns:w="{W}"><w:comment w:author="Writer"/></w:comments>'.encode(),
        "customXml/item1.xml": b"<custom>kept</custom>",
    })
    output = tmp_path / "clean.docx"
    result = clean_docx(marked, output, lambda text: text)
    with ZipFile(output) as archive:
        assert b"rsidR" not in archive.read("word/document.xml")
        assert archive.read("customXml/item1.xml") == b"<custom>kept</custom>"
        assert b'author="Writer"' in archive.read("word/comments.xml")
    assert "comments" in result["container_markers"]
    assert "tracked changes" in result["container_markers"]
    assert "customXml" in result["container_markers"]
    assert result["warnings"]


@pytest.mark.parametrize("name,payload", [
    ("../outside.xml", b"<x/>"),
    ("/outside.xml", b"<x/>"),
    ("word/vbaProject.bin", b"macro"),
    ("word/danger.xml", b'<!DOCTYPE x [<!ENTITY a "xx">]><x>&a;</x>'),
])
def test_rejects_unsafe_packages(sample: Path, tmp_path: Path, name: str, payload: bytes) -> None:
    path = tmp_path / "unsafe.docx"
    rewrite_zip(sample, path, {name: payload})
    with pytest.raises(ValueError):
        inspect_docx(path)


def test_rejects_strict_ooxml(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "strict.docx"
    rewrite_zip(sample, path, {"word/document.xml": b'<w:document xmlns:w="http://purl.oclc.org/ooxml/wordprocessingml/main"><w:body/></w:document>'})
    with pytest.raises(ValueError, match="Strict"):
        extract_docx(path)


def test_rejects_backslash_package_path(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "backslash.docx"
    rewrite_zip(sample, path, {"word/bad.xml": b"<x/>"})
    # Windows' ZipInfo constructor normalizes backslashes, so modify both ZIP
    # header copies to represent an actual hostile archive received from abroad.
    path.write_bytes(path.read_bytes().replace(b"word/bad.xml", b"word\\bad.xml"))
    with pytest.raises(ValueError, match="unsafe package path"):
        inspect_docx(path)


def test_bad_zip_and_same_path_are_rejected(sample: Path, tmp_path: Path) -> None:
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a ZIP")
    with pytest.raises(ValueError):
        inspect_docx(bad)
    with pytest.raises(ValueError):
        clean_docx(sample, sample, lambda text: text)


def test_transform_failure_does_not_create_output(sample: Path, tmp_path: Path) -> None:
    output = tmp_path / "output.docx"
    def fail(text: str) -> str:
        raise ValueError("transform failed")
    with pytest.raises(ValueError, match="transform failed"):
        clean_docx(sample, output, fail)
    assert not output.exists()


@pytest.mark.parametrize("constant,value", [
    ("MAX_COMPRESSED_BYTES", 10),
    ("MAX_EXPANDED_BYTES", 10),
    ("MAX_ENTRIES", 1),
])
def test_resource_limits_are_enforced(sample: Path, monkeypatch, constant: str, value: int) -> None:
    monkeypatch.setattr(docx_io, constant, value)
    with pytest.raises(ValueError, match="limit"):
        inspect_docx(sample)


def test_duplicate_entries_are_rejected(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "duplicates.docx"
    path.write_bytes(sample.read_bytes())
    with pytest.warns(UserWarning, match="Duplicate name"):
        with ZipFile(path, "a") as archive:
            archive.writestr("word/document.xml", b"<x/>")
    with pytest.raises(ValueError, match="duplicate"):
        inspect_docx(path)


def test_malformed_xml_is_rejected(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "malformed.docx"
    rewrite_zip(sample, path, {"word/extra.xml": b"<broken>"})
    with pytest.raises(ValueError, match="Invalid XML"):
        inspect_docx(path)


def test_external_relationship_is_reported_without_access(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "external.docx"
    relationships = b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="remote" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
                    Target="https://example.invalid/never-fetch-this" TargetMode="External"/>
    </Relationships>'''
    rewrite_zip(sample, path, {"word/_rels/extra.xml.rels": relationships})
    assert "external relationships" in inspect_docx(path)["container_markers"]


def test_nonstring_transform_is_rejected(sample: Path, tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="return a string"):
        clean_docx(sample, tmp_path / "output.docx", lambda text: None)


@pytest.mark.parametrize("fragment,marker", [
    ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x=1</m:t></m:r></m:oMath>', "equations"),
    ('<w:sym w:font="Wingdings" w:char="F020"/>', "symbol characters"),
    ('<w:altChunk/>', "alternative-format content"),
    ('<w:fldSimple w:instr="DATE"/>', "field codes"),
    ('<w:footnoteReference w:id="1"/>', "footnotes"),
    ('<w:endnoteReference w:id="1"/>', "endnotes"),
    ('<w:txbxContent><w:p><w:r><w:t>Box text</w:t></w:r></w:p></w:txbxContent>', "text boxes"),
    ('<w:drawing/>', "drawings"),
    ('<w:object/>', "embedded objects"),
    ('<w:ruby/>', "ruby annotations"),
    ('<w:noBreakHyphen/>', "special text elements"),
    ('<w:numPr/>', "numbered lists"),
])
def test_unextracted_content_gets_lossy_marker(sample: Path, tmp_path: Path, fragment: str, marker: str) -> None:
    path = tmp_path / "complex.docx"
    document = f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>Equation:</w:t></w:r>{fragment}</w:p></w:body></w:document>'
    rewrite_zip(sample, path, {"word/document.xml": document.encode()})
    assert marker in inspect_docx(path)["container_markers"]


def test_tables_headers_and_footer_parts_are_reported(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "notes.docx"
    rewrite_zip(sample, path, {
        "word/footnotes.xml": f'<w:footnotes xmlns:w="{W}"/>'.encode(),
        "word/endnotes.xml": f'<w:endnotes xmlns:w="{W}"/>'.encode(),
        "word/footer1.xml": f'<w:ftr xmlns:w="{W}"/>'.encode(),
    })
    markers = set(inspect_docx(path)["container_markers"])
    assert {"tables", "headers", "footers", "footnotes", "endnotes"} <= markers


def test_plain_document_has_no_false_lossy_markers(tmp_path: Path) -> None:
    path = tmp_path / "plain.docx"
    rebuild_docx("正文一行\n第二行", path)
    assert inspect_docx(path)["container_markers"] == []


def test_corrupt_deflate_stream_has_friendly_error(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "bad-deflate.docx"
    with ZipFile(sample) as archive:
        offset = archive.getinfo("word/document.xml").header_offset
    payload = sample.read_bytes()
    filename_length, extra_length = struct.unpack_from("<HH", payload, offset + 26)
    data_offset = offset + 30 + filename_length + extra_length
    path.write_bytes(payload[:data_offset] + b"\xff\xff" + payload[data_offset + 2:])
    with pytest.raises(ValueError, match="Invalid, encrypted, or unsupported"):
        inspect_docx(path)


def test_equation_text_is_included_in_extra_scan(sample: Path, tmp_path: Path) -> None:
    path = tmp_path / "equation.docx"
    xml = f'<w:document xmlns:w="{W}" xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><w:body><w:p><m:oMath><m:r><m:t>x\u200b=1</m:t></m:r></m:oMath></w:p></w:body></w:document>'
    rewrite_zip(sample, path, {"word/document.xml": xml.encode()})
    assert "x\u200b=1" in inspect_docx(path)["extra_text"]


@pytest.mark.parametrize("operation", ["clean", "rebuild"])
def test_existing_destination_is_never_overwritten(sample: Path, tmp_path: Path, operation: str) -> None:
    output = tmp_path / "existing.docx"
    previous = b"existing valuable document bytes"
    output.write_bytes(previous)
    with pytest.raises(FileExistsError):
        if operation == "clean":
            clean_docx(sample, output, lambda text: text)
        else:
            rebuild_docx("new content", output)
    assert output.read_bytes() == previous
